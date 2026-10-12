"""Aragón en línea · Programa del PC de la administración.

Revisa la carpeta "Para subir" y publica lo que encuentre, con las mismas validaciones del panel
(las vuelve a hacer la base de datos: si algo no cuadra, no se publica nada):

  C:\\Aragon en linea\\
    Para subir\\
      Cartera\\                    Excel de cartera (plantilla, con la hoja Control llena)
      Informes\\                   Excel del informe del mes (plantilla, con la hoja Control llena)
      Documentos\\
        Copropietarios\\<Categoría>\\   PDF para propietarios, consejo y administración
        Solo consejo\\<Categoría>\\     PDF solo para el consejo y la administración
    Publicados\\AAAA-MM-DD\\         lo que se publicó (copia de respaldo)
    Con errores\\                    lo que no se publicó, con un archivo "QUE CORREGIR"
    ULTIMO RESULTADO.txt            resumen de la última vez que corrió
    programa\\                       este programa y su archivo .env (clave: NO compartir)

Corre solo con el Programador de tareas de Windows: al prender el PC y cada hora.
Uso manual:  pyw aragon_sync.py           (sin ventana)
             py aragon_sync.py --probar   (comprueba la conexión y la clave)
"""
import datetime as dt
import json
import os
import re
import shutil
import sys
import time
import unicodedata
import uuid
from pathlib import Path

import openpyxl
import requests

import revision_pdf

AQUI = Path(__file__).resolve().parent
VERSION = "1.0"
CATEGORIAS = ["Actas de asamblea", "Actas de consejo", "Estados financieros", "Informes de gestión", "Contratos", "Pólizas", "Otros"]
PERMITIDOS_URL = "https://aragonenlinea.github.io/datos/datos_permitidos.json"
PERMITIDOS_RESPALDO = {"correos": ["aragonenlinea.neiva@gmail.com", "conjuntoaragon2020@gmail.com"],
                       "telefonos": ["3173656199", "3163952209"], "numeros": ["900.140.902"]}
ESPERA_SEGUNDOS = 60      # no toca archivos que se modificaron hace menos de esto (se están copiando)


# ---------------------------------------------------------------- configuración
def leer_env():
    datos = {}
    ruta = AQUI / ".env"
    if ruta.exists():
        for linea in ruta.read_text(encoding="utf-8-sig").splitlines():   # -sig: tolera el BOM del Bloc de notas
            if "=" in linea and not linea.strip().startswith("#"):
                k, v = linea.split("=", 1)
                datos[k.strip()] = v.strip().strip('"')
    return datos


ENV = leer_env()
URL = ENV.get("SUPABASE_URL", "https://bljgybxgkcpnlwlvflss.supabase.co").rstrip("/")
CLAVE = ENV.get("SUPABASE_CLAVE_SECRETA", "")
RAIZ = Path(ENV.get("CARPETA", r"C:\Aragon en linea"))
SUBIR, PUBLICADOS, ERRORES = RAIZ / "Para subir", RAIZ / "Publicados", RAIZ / "Con errores"


def encabezados(extra=None):
    h = {"apikey": CLAVE}
    if CLAVE.startswith("eyJ"):          # clave antigua (JWT service_role): va también en Authorization
        h["Authorization"] = f"Bearer {CLAVE}"
    h.update(extra or {})
    return h


class SinConexion(Exception):
    """Falla de red o del servidor: el archivo se deja donde está y se intenta en la próxima vuelta."""


def rpc(nombre, args):
    try:
        r = requests.post(f"{URL}/rest/v1/rpc/{nombre}", json=args, timeout=60,
                          headers=encabezados({"Content-Type": "application/json"}))
    except requests.RequestException as e:
        raise SinConexion(f"No hay conexión con Supabase ({e.__class__.__name__}).")
    if r.status_code in (401, 403):
        raise SinConexion("Supabase rechazó la clave del programa. El administrador de la página debe revisar el archivo .env.")
    if r.status_code >= 500:
        raise SinConexion(f"Supabase respondió con un error temporal ({r.status_code}).")
    if r.status_code >= 400:
        try:
            m = r.json().get("message", r.text)
        except ValueError:
            m = r.text
        return {"ok": False, "mensaje": f"La base de datos rechazó el archivo: {m}"}
    return r.json()


# ---------------------------------------------------------------- utilidades
def normal(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode().lower().strip()
    return re.sub(r"[^a-z0-9]+", "_", s).strip("_")


def leer_monto(v):
    """1250380 · "1.250.380" · "$ 1.250.380" · "(550.000)" · "-550000" → entero (o None si está vacío)."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        if float(v) != int(v):
            raise ValueError(f"«{v}» tiene decimales")
        return int(v)
    s = str(v).strip()
    negativo = s.startswith("(") and s.endswith(")") or s.startswith("-") or s.startswith("−")
    s = re.sub(r"[$\s()−-]", "", s)
    if re.fullmatch(r"\d{1,3}(\.\d{3})+(,\d+)?|\d+(,\d+)?", s):
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(,\d{3})+(\.\d+)?", s):
        s = s.replace(",", "")
    try:
        n = float(s)
    except ValueError:
        raise ValueError(f"«{v}» no es un número")
    if n != int(n):
        raise ValueError(f"«{v}» tiene decimales")
    return -int(n) if negativo else int(n)


def leer_fecha(v):
    if v in (None, ""):
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        return (v.date() if isinstance(v, dt.datetime) else v).isoformat()
    s = str(v).strip()
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        return f"{m[1]}-{int(m[2]):02d}-{int(m[3]):02d}"
    m = re.match(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})$", s)
    if m:
        return f"{m[3]}-{int(m[2]):02d}-{int(m[1]):02d}"
    raise ValueError(f"fecha «{v}» no válida")


def hoja(wb, nombre, defecto=True):
    for ws in wb.worksheets:
        if normal(ws.title) == nombre:
            return ws
    if defecto:
        for ws in wb.worksheets:
            if normal(ws.title) not in ("instrucciones", "control"):
                return ws
    return None


def filas(ws, alias):
    """Filas de una hoja como diccionarios, con los nombres de columna de la plantilla."""
    it = ws.iter_rows(values_only=True)
    cab = next(it, None) or []
    campos = [alias.get(normal(c)) for c in cab]
    salida = []
    for n, valores in enumerate(it, start=2):
        if not any(v not in (None, "") for v in valores):
            continue
        d = {"_fila": n}
        for campo, v in zip(campos, valores):
            if campo and campo not in d:
                d[campo] = v
        salida.append(d)
    return salida


def control(wb):
    """Hoja Control: columna A el dato, columna B el valor."""
    ws = hoja(wb, "control", defecto=False)
    datos = {}
    if ws:
        for a, b, *_ in ws.iter_rows(values_only=True):
            if a:
                datos[normal(a)] = b
    return datos


# ---------------------------------------------------------------- cartera
ALIAS_CARTERA = {
    "unidad": "unidad", "casa": "unidad", "no_casa": "unidad", "numero_casa": "unidad",
    "fecha_corte": "fecha_corte", "fecha_de_corte": "fecha_corte", "corte": "fecha_corte",
    "periodo": "periodo", "mes": "periodo", "saldo_anterior": "saldo_anterior",
    "cuota_administracion": "cuota_administracion", "cuota_admin": "cuota_administracion", "administracion": "cuota_administracion",
    "cuota_extraordinaria": "cuota_extraordinaria", "cuota_extra": "cuota_extraordinaria", "extraordinaria": "cuota_extraordinaria",
    "parqueadero": "parqueadero", "parqueaderos": "parqueadero", "multas": "multas", "sanciones": "multas", "multas_y_sanciones": "multas",
    "intereses_mora": "intereses_mora", "intereses": "intereses_mora", "intereses_de_mora": "intereses_mora",
    "pagos_periodo": "pagos_periodo", "pagos": "pagos_periodo", "pagos_del_periodo": "pagos_periodo",
    "saldo_total": "saldo_total", "total": "saldo_total", "saldo": "saldo_total",
    "mora_1_30": "mora_1_30", "mora_31_90": "mora_31_90", "mora_91_180": "mora_91_180",
    "mora_181_360": "mora_181_360", "mora_mas_360": "mora_mas_360"}
MONTOS = ["saldo_anterior", "cuota_administracion", "cuota_extraordinaria", "parqueadero", "multas",
          "intereses_mora", "pagos_periodo", "saldo_total", "mora_1_30", "mora_31_90", "mora_91_180", "mora_181_360", "mora_mas_360"]


def procesar_cartera(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    errores, salida, fechas, periodos = [], [], set(), set()
    for f in filas(hoja(wb, "cartera"), ALIAS_CARTERA):
        fila = {}
        casa = re.sub(r"\D", "", str(f.get("unidad") or ""))
        fila["unidad"] = int(casa) if casa else None
        if not fila["unidad"]:
            errores.append(f"Fila {f['_fila']}: falta el número de casa.")
        for k in MONTOS:
            try:
                n = leer_monto(f.get(k))
            except ValueError as e:
                errores.append(f"Fila {f['_fila']} ({k}): {e}.")
                continue
            if n is not None:
                fila[k] = n
        if "saldo_total" not in fila:
            errores.append(f"Fila {f['_fila']}: falta el saldo total (si usa fórmulas, abra el archivo en Excel y guárdelo).")
        try:
            fc = leer_fecha(f.get("fecha_corte"))
            if fc:
                fechas.add(fc)
        except ValueError as e:
            errores.append(f"Fila {f['_fila']}: {e}.")
        if f.get("periodo"):
            periodos.add(str(f["periodo"]).strip())
        salida.append(fila)
    c = control(wb)
    try:
        total = leer_monto(c.get("total_informe", c.get("total_cartera_informe")))
        if total is None:
            errores.append("Llene la hoja Control con el total de la cartera según el informe contable (total_informe).")
    except ValueError as e:
        total = None
        errores.append(f"Hoja Control: el total {e}.")
    if len(fechas) != 1:
        errores.append("Todas las filas deben tener la misma fecha de corte." if fechas else "Falta la fecha de corte.")
    if len(periodos) != 1:
        errores.append("Todas las filas deben tener el mismo periodo." if periodos else "Falta el periodo (por ejemplo: Septiembre 2026).")
    if errores:
        return {"ok": False, "mensaje": "No se publicó.", "errores": errores}
    return rpc("importar_cartera", {"p_fecha_corte": fechas.pop(), "p_periodo": periodos.pop(), "p_total_informe": total,
                                    "p_archivo": ruta.name[:200], "p_filas": salida, "p_origen": "script"})


# ---------------------------------------------------------------- informe mensual
ALIAS_EJEC = {"tipo": "tipo", "rubro": "rubro", "cuenta": "rubro", "valor": "valor", "ejecutado": "valor", "valor_mes": "valor"}
ALIAS_PAGO = {"fecha": "fecha", "rubro": "rubro", "beneficiario": "beneficiario", "tercero": "beneficiario", "proveedor": "beneficiario",
              "concepto": "concepto", "detalle": "concepto", "descripcion": "concepto", "valor": "valor"}


def leer_tipo(v):
    n = normal(v)
    return "ingreso" if n.startswith("ingreso") else "gasto" if n.startswith(("gasto", "egreso")) else None


def procesar_informe(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    errores, ejec, pagos = [], [], []
    for f in filas(hoja(wb, "ejecucion"), ALIAS_EJEC):
        tipo = leer_tipo(f.get("tipo"))
        if not tipo:
            errores.append(f"Ejecución, fila {f['_fila']}: el tipo debe ser «ingreso» o «gasto».")
        try:
            valor = leer_monto(f.get("valor")) or 0
        except ValueError as e:
            errores.append(f"Ejecución, fila {f['_fila']}: {e}.")
            valor = 0
        ejec.append({"tipo": tipo, "rubro": str(f.get("rubro") or "").strip(), "valor": valor})
    ws_pagos = hoja(wb, "pagos", defecto=False)
    if ws_pagos:
        for f in filas(ws_pagos, ALIAS_PAGO):
            try:
                valor = leer_monto(f.get("valor"))
                fecha = leer_fecha(f.get("fecha")) or ""
            except ValueError as e:
                errores.append(f"Pagos, fila {f['_fila']}: {e}.")
                continue
            pagos.append({"fecha": fecha, "rubro": str(f.get("rubro") or "").strip(), "beneficiario": str(f.get("beneficiario") or "").strip(),
                          "concepto": str(f.get("concepto") or "").strip()[:200], "valor": valor})
    c = control(wb)
    m = re.search(r"(20\d\d)[_\- ](\d{1,2})(?!\d)", ruta.name)
    try:
        anio = int(c.get("anio") or c.get("ano") or (m and m[1]) or 0)
        mes = int(c.get("mes") or (m and m[2]) or 0)
        ti, tg = leer_monto(c.get("total_ingresos")), leer_monto(c.get("total_gastos"))
        tp = leer_monto(c.get("total_pagos")) if pagos else None
    except ValueError as e:
        return {"ok": False, "mensaje": "No se publicó.", "errores": [f"Hoja Control: {e}."]}
    if not (2020 <= anio <= 2100 and 1 <= mes <= 12):
        errores.append("Escriba el año y el mes en la hoja Control.")
    if ti is None or tg is None:
        errores.append("Escriba en la hoja Control los totales de ingresos y gastos del informe contable.")
    if pagos and tp is None:
        errores.append("El archivo trae pagos: escriba en la hoja Control el total de pagos del mes.")
    if errores:
        return {"ok": False, "mensaje": "No se publicó.", "errores": errores}
    return rpc("importar_informe_mes", {"p_anio": anio, "p_mes": mes, "p_total_ingresos": ti, "p_total_gastos": tg,
                                        "p_total_pagos": tp, "p_archivo": ruta.name[:200], "p_ejecucion": ejec,
                                        "p_pagos": pagos, "p_origen": "script"})


# ---------------------------------------------------------------- documentos
def permitidos():
    try:
        r = requests.get(PERMITIDOS_URL, timeout=20)
        if r.ok:
            return r.json()
    except (requests.RequestException, ValueError):
        pass
    return PERMITIDOS_RESPALDO


def procesar_documento(ruta, audiencia, categoria):
    # Nombre del archivo: "2026-03-14 Acta de asamblea ordinaria.pdf" (la fecha es la del documento).
    m = re.match(r"^(\d{4}-\d{2}-\d{2})\s*[-_ ]\s*(.+)$", ruta.stem)
    if m:
        fecha, titulo = m[1], m[2].strip()
    else:
        fecha, titulo = dt.date.fromtimestamp(ruta.stat().st_mtime).isoformat(), ruta.stem.strip()
    titulo = re.sub(r"\s+", " ", titulo.replace("_", " "))[:150]
    try:
        datos, paginas, hallazgos, imagenes, _ = revision_pdf.revisar(ruta, permitidos())
    except ValueError as e:
        return {"ok": False, "mensaje": str(e)}
    problemas = []
    if hallazgos:
        problemas.append("Tiene posibles datos personales: " + "; ".join(hallazgos[:15]) + ("…" if len(hallazgos) > 15 else "."))
    if imagenes:
        problemas.append(f"Las páginas {', '.join(map(str, imagenes))} son imagen (escaneadas) y no se pueden revisar solas.")
    if problemas:
        problemas.append("Súbalo desde el panel (Documentos): allí puede revisarlo, usar «Tapar datos» y confirmarlo.")
        return {"ok": False, "mensaje": "No se publicó.", "errores": problemas}
    if len(datos) > 50 * 1024 * 1024:
        return {"ok": False, "mensaje": "El PDF pesa más de 50 MB. Redúzcalo y vuelva a ponerlo en la carpeta."}
    destino = f"docs/{dt.date.today().year}/{uuid.uuid4()}.pdf"
    try:
        r = requests.post(f"{URL}/storage/v1/object/privados/{destino}", data=datos, timeout=120,
                          headers=encabezados({"Content-Type": "application/pdf", "x-upsert": "false"}))
    except requests.RequestException as e:
        raise SinConexion(f"No hay conexión con Supabase ({e.__class__.__name__}).")
    if r.status_code in (401, 403):
        raise SinConexion("Supabase rechazó la clave del programa al subir el PDF.")
    if not r.ok:
        raise SinConexion(f"No se pudo subir el PDF ({r.status_code}): {r.text[:200]}")
    res = rpc("registrar_documento", {"p_titulo": titulo, "p_categoria": categoria, "p_fecha": fecha, "p_descripcion": None,
                                      "p_audiencia": audiencia, "p_ruta": destino, "p_tamano": len(datos), "p_paginas": paginas,
                                      "p_revisado": True, "p_origen": "script", "p_hallazgos": 0})
    if not res.get("ok"):
        try:
            requests.delete(f"{URL}/storage/v1/object/privados/{destino}", headers=encabezados(), timeout=60)
        except requests.RequestException:
            pass
    return res


def categoria_de(carpeta):
    n = normal(carpeta)
    for c in CATEGORIAS:
        if normal(c) == n:
            return c
    return "Otros"


# ---------------------------------------------------------------- recorrido de carpetas
def preparar_carpetas():
    for c in [SUBIR / "Cartera", SUBIR / "Informes", PUBLICADOS, ERRORES] + \
             [SUBIR / "Documentos" / a / cat for a in ("Copropietarios", "Solo consejo") for cat in CATEGORIAS]:
        c.mkdir(parents=True, exist_ok=True)


def listo(ruta):
    nombre = ruta.name
    if nombre.startswith(("~$", ".")) or nombre.lower() in ("desktop.ini", "thumbs.db"):
        return False
    return time.time() - ruta.stat().st_mtime >= ESPERA_SEGUNDOS


def mover(ruta, carpeta):
    carpeta.mkdir(parents=True, exist_ok=True)
    destino = carpeta / ruta.name
    if destino.exists():
        destino = carpeta / f"{ruta.stem} ({dt.datetime.now():%H%M%S}){ruta.suffix}"
    shutil.move(str(ruta), str(destino))
    return destino


def pendientes():
    for ruta in sorted((SUBIR / "Cartera").glob("*")):
        if ruta.is_file():
            yield ruta, "Cartera", lambda r: procesar_cartera(r) if r.suffix.lower() in (".xlsx", ".xlsm") else \
                {"ok": False, "mensaje": "Use la plantilla de cartera en Excel (.xlsx)."}
    for ruta in sorted((SUBIR / "Informes").glob("*")):
        if ruta.is_file():
            yield ruta, "Informe", lambda r: procesar_informe(r) if r.suffix.lower() in (".xlsx", ".xlsm") else \
                {"ok": False, "mensaje": "Use la plantilla de informe mensual en Excel (.xlsx)."}
    for audiencia, carpeta in (("copropietarios", "Copropietarios"), ("consejo", "Solo consejo")):
        base = SUBIR / "Documentos" / carpeta
        for ruta in sorted(base.rglob("*")):
            if ruta.is_file():
                cat = categoria_de(ruta.parent.name) if ruta.parent != base else "Otros"
                yield ruta, f"Documento ({carpeta} · {cat})", (lambda a, c: lambda r: procesar_documento(r, a, c) if r.suffix.lower() == ".pdf"
                                                              else {"ok": False, "mensaje": "Solo se suben archivos PDF."})(audiencia, cat)


def registrar(lineas):
    archivo = RAIZ / "registro.txt"
    anteriores = archivo.read_text(encoding="utf-8").splitlines()[-3000:] if archivo.exists() else []
    archivo.write_text("\n".join(anteriores + lineas) + "\n", encoding="utf-8")


def ejecutar():
    preparar_carpetas()
    candado = RAIZ / "programa" / ".corriendo"
    if candado.exists() and time.time() - candado.stat().st_mtime < 3600:
        return 0            # ya hay otra vuelta en curso
    candado.parent.mkdir(parents=True, exist_ok=True)
    candado.write_text(str(os.getpid()))
    ahora = dt.datetime.now()
    lineas, resumen = [], []
    try:
        if not CLAVE:
            resumen.append("Falta la clave del programa (.env). El administrador de la página debe correr el instalador.")
        else:
            for ruta, tipo, proceso in pendientes():
                if not listo(ruta):
                    continue
                try:
                    res = proceso(ruta)
                except SinConexion as e:
                    resumen.append(f"PENDIENTE  {ruta.name}: {e} Se intentará de nuevo.")
                    lineas.append(f"{ahora:%Y-%m-%d %H:%M} PENDIENTE {tipo} {ruta.name}: {e}")
                    break   # sin conexión: no tiene sentido seguir con los demás
                except Exception as e:      # nada debe quedar a medias por un error inesperado
                    res = {"ok": False, "mensaje": f"Error inesperado al leer el archivo: {e.__class__.__name__}: {e}"}
                if res.get("ok"):
                    mover(ruta, PUBLICADOS / f"{ahora:%Y-%m-%d}")
                    resumen.append(f"PUBLICADO  {ruta.name}: {res.get('mensaje', '')}")
                    lineas.append(f"{ahora:%Y-%m-%d %H:%M} PUBLICADO {tipo} {ruta.name}: {res.get('mensaje', '')}")
                else:
                    destino = mover(ruta, ERRORES)
                    detalle = [res.get("mensaje", "No se publicó.")] + [f"  - {e}" for e in (res.get("errores") or [])]
                    (ERRORES / f"{destino.stem} - QUE CORREGIR.txt").write_text(
                        f"Archivo: {destino.name}\nTipo: {tipo}\nFecha: {ahora:%Y-%m-%d %H:%M}\n\n" + "\n".join(detalle) +
                        "\n\nCorrija el archivo y vuelva a ponerlo en la carpeta «Para subir». Puede borrar este aviso.\n", encoding="utf-8")
                    resumen.append(f"CON ERRORES  {ruta.name}: ver «Con errores»")
                    lineas.append(f"{ahora:%Y-%m-%d %H:%M} ERROR {tipo} {ruta.name}: " + " | ".join(detalle))
        texto = f"Aragón en línea · última revisión: {ahora:%Y-%m-%d %H:%M}\n\n" + ("\n".join(resumen) if resumen else "No había archivos nuevos.") + "\n"
        (RAIZ / "ULTIMO RESULTADO.txt").write_text(texto, encoding="utf-8")
        if lineas:
            registrar(lineas)
    finally:
        candado.unlink(missing_ok=True)
    return 0


def probar():
    print(f"Aragón en línea · programa {VERSION}")
    print(f"Carpeta: {RAIZ}")
    if not CLAVE:
        print("FALTA la clave en el archivo .env.")
        return 1
    try:
        r = requests.get(f"{URL}/storage/v1/bucket/privados", headers=encabezados(), timeout=30)
    except requests.RequestException as e:
        print(f"Sin conexión con Supabase: {e}")
        return 1
    if r.ok:
        print("Conexión y clave correctas: el programa puede publicar.")
        return 0
    print(f"Supabase rechazó la clave ({r.status_code}). Revise que copió la clave SECRETA (sb_secret_...), completa.")
    return 1


if __name__ == "__main__":
    sys.exit(probar() if "--probar" in sys.argv else ejecutar())
